"""Tests for document upload and labeling-template generation."""

from __future__ import annotations

import asyncio
import json
from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from fpdf import FPDF

from legal_eval_api.dataset_sources.documents import (
    MAX_FILE_BYTES,
    TARGET_MAX_CHARS,
    build_candidates_from_uploads,
    chunk_page_sections,
    chunk_paragraph_sections,
    excerpts_from_upload,
    extract_docx_text,
    extract_pdf_text,
    render_template_jsonl,
    template_row,
)
from legal_eval_api.db import init_db
from legal_eval_api.main import app
from legal_eval_api.orgs import register_org
from legal_eval_api.schemas import CreateOrgRequest


def _make_pdf_bytes(*pages: str) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=False)
    for text in pages:
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.multi_cell(0, 8, text)
    return bytes(pdf.output())


def _make_docx_bytes(*paragraphs: str) -> bytes:
    document = Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def api_client(tmp_path, monkeypatch) -> TestClient:
    monkeypatch.setenv("LEGAL_EVAL_API_DATA", str(tmp_path / "data"))
    init_db()
    return TestClient(app)


@pytest.fixture
def org_headers(api_client: TestClient) -> dict[str, str]:
    created = register_org(CreateOrgRequest(name="Doc Test Org"))
    return {"Authorization": f"Bearer {created.api_key}"}


def test_extract_pdf_text_from_synthetic_pdf() -> None:
    data = _make_pdf_bytes(
        "Payment is due within thirty (30) days of invoice date.",
        "Either party may terminate with thirty days written notice.",
    )
    pages = extract_pdf_text(data, filename="msa.pdf")
    assert len(pages) == 2
    assert "Payment is due" in pages[0][1]
    assert pages[0][0] == "page 1"


def test_extract_docx_text() -> None:
    data = _make_docx_bytes(
        "Confidentiality obligations survive termination.",
        "Governing law is Delaware.",
    )
    paragraphs = extract_docx_text(data, filename="nda.docx")
    assert len(paragraphs) == 2
    assert paragraphs[0][1].startswith("Confidentiality")


def test_chunk_page_sections_merges_short_pages() -> None:
    short = "A" * 250
    chunks = chunk_page_sections([("page 1", short), ("page 2", short)])
    assert len(chunks) == 1
    assert len(chunks[0][1]) >= 500


def test_chunk_paragraph_sections_respects_target_size() -> None:
    paragraphs = [(f"paragraph {index}", "Clause text. " * 120) for index in range(1, 6)]
    chunks = chunk_paragraph_sections(paragraphs)
    assert len(chunks) >= 2
    assert all(len(text) <= TARGET_MAX_CHARS for _, text in chunks)
    assert 400 <= len(chunks[0][1]) <= TARGET_MAX_CHARS


def test_excerpts_from_upload_txt() -> None:
    text = "\n\n".join(f"Section {index}: " + ("Lorem ipsum. " * 80) for index in range(1, 4))
    excerpts = excerpts_from_upload("contract.txt", text.encode("utf-8"))
    assert len(excerpts) >= 1
    assert excerpts[0].contract_title == "contract"
    assert excerpts[0].contract_excerpt


def test_template_row_has_empty_label_placeholders() -> None:
    from legal_eval_api.dataset_sources.documents import CandidateExcerpt

    row = template_row(
        CandidateExcerpt(
            id="contract-001",
            contract_excerpt="Sample clause text.",
            contract_title="MSA",
            source_filename="msa.pdf",
            source_reference="page 1",
        ),
    )
    assert row["category"] == ""
    assert row["present"] is None
    assert row["gold_spans"] == []


def test_render_template_jsonl_is_parseable() -> None:
    from legal_eval_api.dataset_sources.documents import CandidateExcerpt

    candidates = [
        CandidateExcerpt(
            id="contract-001",
            contract_excerpt="Payment due in 30 days.",
            contract_title="MSA",
            source_filename="msa.pdf",
            source_reference="page 1",
        ),
    ]
    rendered = render_template_jsonl(candidates)
    rows = [json.loads(line) for line in rendered.strip().splitlines()]
    assert len(rows) == 1
    assert set(rows[0].keys()) == {
        "id",
        "contract_excerpt",
        "contract_title",
        "category",
        "present",
        "gold_spans",
    }
    assert rows[0]["present"] is None
    assert rows[0]["category"] == ""


def test_build_candidates_rejects_unsupported_type() -> None:
    from fastapi import HTTPException, UploadFile

    upload = UploadFile(filename="data.csv", file=BytesIO(b"a,b,c"))
    with pytest.raises(HTTPException) as exc:
        asyncio.run(build_candidates_from_uploads([upload]))
    assert exc.value.status_code == 400
    assert "Unsupported file type" in exc.value.detail


def test_build_candidates_rejects_oversized_file() -> None:
    from fastapi import HTTPException, UploadFile

    upload = UploadFile(
        filename="big.txt",
        file=BytesIO(b"x" * (MAX_FILE_BYTES + 1)),
    )
    with pytest.raises(HTTPException) as exc:
        asyncio.run(build_candidates_from_uploads([upload]))
    assert exc.value.status_code == 400
    assert "MiB limit" in exc.value.detail


def test_excerpts_from_empty_pdf_raises() -> None:
    from fastapi import HTTPException

    pdf = FPDF()
    pdf.add_page()
    with pytest.raises(HTTPException) as exc:
        excerpts_from_upload("blank.pdf", bytes(pdf.output()))
    assert exc.value.status_code == 400
    assert "No extractable text" in exc.value.detail


def test_api_import_documents_returns_staging_and_template(
    api_client: TestClient,
    org_headers: dict[str, str],
) -> None:
    pdf_bytes = _make_pdf_bytes("Indemnification cap is limited to fees paid in the prior year.")
    response = api_client.post(
        "/api/v1/datasets/import/documents",
        headers=org_headers,
        files=[("files", ("msa.pdf", pdf_bytes, "application/pdf"))],
    )
    assert response.status_code == 200, response.text
    payload = response.json()
    assert payload["excerpt_count"] >= 1
    assert payload["staging_id"]
    assert payload["preview"]
    assert "two-step workflow" in payload["instructions"]

    download = api_client.get(payload["download_path"], headers=org_headers)
    assert download.status_code == 200
    rows = [json.loads(line) for line in download.text.strip().splitlines()]
    assert len(rows) == payload["excerpt_count"]
    assert rows[0]["present"] is None
    assert rows[0]["category"] == ""
    assert rows[0]["gold_spans"] == []


def test_api_import_documents_docx(api_client: TestClient, org_headers: dict[str, str]) -> None:
    docx_bytes = _make_docx_bytes("Assignment requires prior written consent.")
    response = api_client.post(
        "/api/v1/datasets/import/documents",
        headers=org_headers,
        files=[
            (
                "files",
                ("assignment.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            ),
        ],
    )
    assert response.status_code == 200
    assert response.json()["excerpt_count"] >= 1


def test_api_import_documents_bad_file_type(
    api_client: TestClient,
    org_headers: dict[str, str],
) -> None:
    response = api_client.post(
        "/api/v1/datasets/import/documents",
        headers=org_headers,
        files=[("files", ("notes.csv", b"a,b,c", "text/csv"))],
    )
    assert response.status_code == 400
    assert "Unsupported file type" in response.json()["detail"]


def test_api_download_template_requires_auth(api_client: TestClient) -> None:
    response = api_client.get("/api/v1/datasets/import/documents/missing/template.jsonl")
    assert response.status_code == 401
